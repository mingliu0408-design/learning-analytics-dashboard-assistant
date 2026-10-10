#!/usr/bin/env python3
"""Bounded local profiling of E0 exact locators. No rows or identities exported."""
from pathlib import Path
from collections import Counter
import csv
import hashlib
import json

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
E0 = HERE.parent / "e0_crosspage_functional_check"
LEGACY = Path("/Users/liuming/Downloads/duteaiapp_cfull_semantic_integrated_v1_副本3")


def count_label(n):
    return n if n == 0 or n >= 10 else "SUPPRESSED_LT_10"


def profile():
    manifest = json.loads((E0 / "SOURCE_MANIFEST.json").read_text())
    expected = {r["path"]: r for r in manifest["raw_source_metadata_only"]}
    entries = []
    for row in csv.DictReader((E0 / "RAW_DASHBOARD_SOURCE_MAP.csv").open()):
        if row["location"] == "optional_API_only":
            entries.append({"binding_field": row["binding_field"], "status": "NOT_BOUND", "rows_exported": 0})
            continue
        path = Path(row["path"])
        if not (path.is_relative_to(LEGACY) or path.is_relative_to(ROOT.parent)):
            raise RuntimeError("Locator outside authorised roots")
        old = expected[str(path)]
        record = {"binding_field": row["binding_field"], "location": row["location"],
                  "path": str(path), "rows_exported": 0, "actual_local_profiling": False,
                  "course_binding": "UNVERIFIED_NO_OPERATIONAL_COURSE_MANIFEST",
                  "coverage_completeness": "UNKNOWN", "source_kind": "ORIGINAL_LOCAL_SOURCE"}
        if not path.is_file():
            record["status"] = "SOURCE_DATA_UNAVAILABLE_AT_CURRENT_PATH"
            entries.append(record)
            continue
        data = path.read_bytes()
        sha = hashlib.sha256(data).hexdigest()
        record.update(sha256=sha, bytes=len(data), matches_E0_bytes=sha == old["sha256"])
        if sha != old["sha256"]:
            record["status"] = "SOURCE_HASH_CHANGED_NOT_PROFILED"
            entries.append(record)
            continue
        if path.suffix == ".pkl":
            # E0 explicitly records no source commit/provenance claim. A byte
            # match to an opaque pickle is not evidence its deserialiser is safe.
            record.update(status="BINARY_SOURCE_NOT_SAFELY_READABLE",
                          schema_authority="COMMITTED_CODE_CONTRACT_ONLY",
                          required_export="CSV or Parquet with original source hash, schema and observation semantics",
                          reason="E0 hashes bytes but records no trusted pickle creation provenance; no unpickle attempted.")
            entries.append(record)
            continue
        if path.suffix == ".csv":
            handle = path.open(newline="", encoding="utf-8-sig")
            reader = csv.reader(handle)
            headers = next(reader)
            rows = reader
        elif path.suffix == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            rows = book.active.iter_rows(values_only=True)
            headers = [str(v) if v is not None else "UNNAMED_COLUMN" for v in next(rows)]
            handle = book
        else:
            raise RuntimeError("Unexpected source format")
        count = 0
        missing = Counter()
        numeric = Counter()
        for values in rows:
            count += 1
            for i, field in enumerate(headers):
                value = values[i] if i < len(values) else None
                if value is None or str(value).strip() == "":
                    missing[field] += 1
                else:
                    try:
                        float(value)
                        numeric[field] += 1
                    except (ValueError, TypeError):
                        pass
        handle.close()
        record.update(status="ACTUAL_LOCAL_PROFILE_COMPLETED", actual_local_profiling=True,
                      verified_headers=headers, row_count=count_label(count),
                      schema_authority="VERIFIED_HEADER_AND_LOCAL_ROWS",
                      missingness={h: "ALL_MISSING" if missing[h] == count else "ALL_PRESENT" if missing[h] == 0 else "SOME_MISSING_COUNTS_SUPPRESSED" for h in headers},
                      field_types={h: "NUMERIC" if numeric[h] == count - missing[h] and numeric[h] else "TEXT_OR_MIXED" for h in headers},
                      course_column_present=any(h.casefold() in {"course", "course_id", "courseid"} for h in headers),
                      time_fields=[h for h in headers if h in {"Time", "Week", "week", "Content_week", "day", "dayOfWeek", "Day_of_week"}])
        if row["binding_field"] == "echo":
            record["temporal_semantics"] = "Content_week is material week; watching time unavailable; values cumulative by recording."
        elif row["binding_field"] == "features":
            record["temporal_semantics"] = "Aggregate features have no interval fields; cannot override scoped event totals."
        elif row["binding_field"] == "gdocs_ts":
            record["temporal_semantics"] = "Cohort day series; not same-learner evidence; not bound in E1."
        else:
            record["temporal_semantics"] = row["time_scope"]
        entries.append(record)
    # E1's own expected locations are checked explicitly, without discovering
    # files in M0/P0/R1 worktrees or silently binding the known legacy sources.
    e1_locations = []
    for r in entries:
        if r.get("location") == "current_checkout":
            rel = Path(r["path"]).relative_to(ROOT.parent / "rq3-rebuild-e0-crosspage")
            e1_locations.append({"path": str(ROOT / rel), "exists": (ROOT / rel).is_file(), "binding_field": r["binding_field"]})
    result = {"task": "RQ3-REBUILD-E1", "sources": entries, "E1_expected_paths": e1_locations,
              "privacy": "No row samples, private values, free text, identities or stratified/small-cell aggregates exported. Counts 1-9 suppressed.",
              "real_experiment_ready": False, "live_configuration": "LIVE_CONFIGURATION_UNCONFIRMED",
              "real_source_aggregation": "NOT_RUN; profiling only; synthetic acceptance has separate fixtures"}
    (HERE / "DATA_COVERAGE_PROFILE.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"profiled": sum(r.get("actual_local_profiling", False) for r in entries),
                      "status_counts": dict(Counter(r["status"] for r in entries)), "private_rows_exported": 0}))


if __name__ == "__main__":
    profile()
